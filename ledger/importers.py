from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook, load_workbook

MAX_FILE_MB = 5
MAX_ROWS = 5000
ACCOUNT_HEADERS = ["code", "name", "type"]
ENTRY_HEADERS = [
    "voucher_no",
    "date",
    "reference",
    "description",
    "account_code",
    "debit",
    "credit",
    "narration",
]
DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y")


class ExcelFormatError(Exception):
    """The file itself is unusable (wrong type, missing columns, too big)."""


# ---------- reading the file ----------
def get_uploaded_file(request):
    f = request.FILES.get("file")
    if not f:
        raise ExcelFormatError('Attach an .xlsx file in the form field named "file".')
    if not f.name.lower().endswith(".xlsx"):
        raise ExcelFormatError("Only .xlsx files are supported.")
    if f.size > MAX_FILE_MB * 1024 * 1024:
        raise ExcelFormatError(f"File is larger than {MAX_FILE_MB} MB.")
    return f


def _norm_header(value):
    return str(value).strip().lower().replace(" ", "_") if value is not None else ""


def read_sheet(file, required_headers):
    """Returns a list of (excel_row_number, {header: value}), skipping blank rows."""
    try:
        wb = load_workbook(file, read_only=True, data_only=True)
    except Exception:
        raise ExcelFormatError("Could not read the file. Upload a valid .xlsx file.")

    try:
        rows = wb.active.iter_rows(values_only=True)
        try:
            header = [_norm_header(h) for h in next(rows)]
        except StopIteration:
            raise ExcelFormatError("The sheet is empty.")

        missing = [h for h in required_headers if h not in header]
        if missing:
            raise ExcelFormatError(f"Missing columns: {', '.join(missing)}")

        result = []
        for n, row in enumerate(rows, start=2):  # row 1 is the header
            if all(c is None or str(c).strip() == "" for c in row):
                continue
            result.append((n, dict(zip(header, row))))
            if len(result) > MAX_ROWS:
                raise ExcelFormatError(f"Too many rows (limit is {MAX_ROWS}).")
        if not result:
            raise ExcelFormatError("No data rows found under the header.")
        return result
    finally:
        wb.close()


# ---------- cell helpers ----------
def as_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():  # Excel gives 1001 as 1001.0
        value = int(value)
    return str(value).strip()


def as_amount(value, label, problems):
    if value is None or str(value).strip() == "":
        return Decimal("0")
    try:
        amount = Decimal(str(value).replace(",", "").strip())
    except InvalidOperation:
        problems.append(f"{label} is not a number")
        return Decimal("0")
    if amount != amount.quantize(Decimal("0.01")):
        problems.append(f"{label} has more than 2 decimal places")
    if amount < 0:
        problems.append(f"{label} cannot be negative")
    return amount


def as_date(value, problems):
    """Returns 'YYYY-MM-DD' or None when the cell is blank."""
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(str(value).strip(), fmt).date().isoformat()
        except ValueError:
            continue
    problems.append(f"date '{value}' is not valid (use YYYY-MM-DD or DD/MM/YYYY)")
    return None


# ---------- parsing ----------
def parse_accounts(rows):
    items, item_rows, errors = [], [], []
    for n, r in rows:
        code = as_text(r.get("code"))
        name = as_text(r.get("name"))
        typ = as_text(r.get("type")).lower()
        problems = [
            f"{label} is required"
            for label, val in (("code", code), ("name", name), ("type", typ))
            if not val
        ]
        if problems:
            errors.append({"row": n, "errors": problems})
            continue
        items.append({"code": code, "name": name, "type": typ})
        item_rows.append(n)
    return items, item_rows, errors


def parse_entries(rows, code_to_id):
    groups, errors = {}, []
    for n, r in rows:
        problems = []
        key = as_text(r.get("voucher_no"))
        if not key:
            errors.append({"row": n, "errors": ["voucher_no is required"]})
            continue

        code = as_text(r.get("account_code"))
        account_id = code_to_id.get(code)
        if not code:
            problems.append("account_code is required")
        elif account_id is None:
            problems.append(f"account code '{code}' not found in this book")

        debit = as_amount(r.get("debit"), "debit", problems)
        credit = as_amount(r.get("credit"), "credit", problems)
        row_date = as_date(r.get("date"), problems)

        g = groups.setdefault(
            key,
            {
                "voucher": key,
                "rows": [],
                "date": None,
                "reference": "",
                "description": "",
                "lines": [],
            },
        )
        g["rows"].append(n)

        if row_date:
            if g["date"] and g["date"] != row_date:
                problems.append("date differs from the first row of this voucher")
            g["date"] = g["date"] or row_date
        g["reference"] = g["reference"] or as_text(r.get("reference"))
        g["description"] = g["description"] or as_text(r.get("description"))

        if problems:
            errors.append({"row": n, "errors": problems})
            continue
        g["lines"].append(
            {
                "account": account_id,
                "debit": str(debit),
                "credit": str(credit),
                "narration": as_text(r.get("narration")),
            }
        )

    for g in groups.values():
        missing = [f for f in ("date", "description") if not g[f]]
        if missing:
            errors.append(
                {
                    "voucher": g["voucher"],
                    "rows": g["rows"],
                    "errors": [
                        f"{', '.join(missing)} required on the first row of the voucher"
                    ],
                }
            )

    entries = [
        {
            "voucher": g["voucher"],
            "rows": g["rows"],
            "data": {
                "date": g["date"],
                "reference": g["reference"],
                "description": g["description"],
                "lines": g["lines"],
            },
        }
        for g in groups.values()
    ]
    return entries, errors


# ---------- downloadable templates ----------
def build_template(kind):
    wb = Workbook()
    ws = wb.active
    if kind == "accounts":
        ws.title = "Accounts"
        ws.append(ACCOUNT_HEADERS)
        ws.append(["1001", "Cash in Hand", "asset"])
        ws.append(["3001", "Owner's Capital", "equity"])
        ws.append(["4001", "Sales", "income"])
    else:
        ws.title = "Journal"
        ws.append(ENTRY_HEADERS)
        ws.append(
            [1, "2026-09-01", "V-001", "Capital introduced", "1001", 500000, None, None]
        )
        ws.append([1, None, None, None, "3001", None, 500000, None])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
