
"""CSV and XLSX exports for journals and account ledgers."""

import csv
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import Account, JournalEntry, JournalLine


MONEY_FORMAT = '#,##0.00;[Red](#,##0.00)'
ZERO = Decimal("0.00")


def _format(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def _safe_csv(value):
    value = _format(value)
    if isinstance(value, str) and value.lstrip().startswith(
        ("=", "+", "-", "@", "\t", "\r", "\n")
    ):
        return "'" + value
    return value


def _response_csv(headers, rows, filename):
    response = HttpResponse(
        content_type="text/csv; charset=utf-8"
    )
    response["Content-Disposition"] = (
        f'attachment; filename="{filename}"'
    )
    response.write("\ufeff")

    writer = csv.writer(response)
    writer.writerow(headers)

    for row in rows:
        writer.writerow([_safe_csv(value) for value in row])

    return response


def _new_sheet(wb, title, headers):
    ws = wb.create_sheet(title=title)
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill(
            "solid", fgColor="243B53"
        )
        cell.alignment = Alignment(
            vertical="center"
        )

    ws.row_dimensions[1].height = 25
    ws.freeze_panes = "A2"
    return ws


def _fill_sheet(ws, rows, money_columns):
    for row in rows:
        ws.append([_format(value) for value in row])

    for col in money_columns:
        for cells in ws.iter_cols(
            min_col=col,
            max_col=col,
            min_row=2,
        ):
            for cell in cells:
                cell.number_format = MONEY_FORMAT

    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        max_length = max(
            (
                len(str(cell.value or ""))
                for cell in column
            ),
            default=10,
        )
        ws.column_dimensions[letter].width = min(
            max(max_length + 3, 14), 38
        )

    ws.auto_filter.ref = ws.dimensions


def _response_xlsx(sheets, filename):
    wb = Workbook()
    wb.remove(wb.active)

    for title, headers, rows, money_columns in sheets:
        ws = _new_sheet(wb, title, headers)
        _fill_sheet(ws, rows, money_columns)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        f'attachment; filename="{filename}"'
    )
    return response


def _deliver(sheets, fmt, filename):
    if fmt == "csv":
        headers = sheets[0][1]
        rows = (
            row
            for _, _, sheet_rows, _ in sheets
            for row in sheet_rows
        )
        return _response_csv(
            headers, rows, f"{filename}.csv"
        )

    return _response_xlsx(
        sheets, f"{filename}.xlsx"
    )


JOURNAL_HEADERS = [
    "Entry No",
    "Date",
    "Reference",
    "Description",
    "Account Code",
    "Account Name",
    "Narration",
    "Debit",
    "Credit",
]


def export_journals(book, fmt, date_from=None, date_to=None):
    entries = JournalEntry.objects.filter(book=book)

    if date_from:
        entries = entries.filter(date__gte=date_from)
    if date_to:
        entries = entries.filter(date__lte=date_to)

    entries = entries.order_by("date", "entry_no")

    lines = (
        JournalLine.objects
        .filter(entry__in=entries)
        .select_related("entry", "account")
        .order_by("entry__date", "entry__entry_no", "pk")
    )

    rows = [
        [
            line.entry.entry_no,
            line.entry.date,
            line.entry.reference,
            line.entry.description,
            line.account.code,
            line.account.name,
            line.narration,
            line.debit,
            line.credit,
        ]
        for line in lines
    ]

    sheets = [
        ("Journal Entries", JOURNAL_HEADERS, rows, [8, 9])
    ]

    return _deliver(
        sheets, fmt, f"book_{book.pk}_journal_entries"
    )


LEDGER_HEADERS = [
    "Account Code",
    "Account Name",
    "Date",
    "Entry No",
    "Reference",
    "Details",
    "Debit",
    "Credit",
    "Balance",
    "Balance Side",
]


def _ledger_rows(book, account, date_from=None, date_to=None):
    # Include older transactions to calculate opening balance.
    lines = (
        JournalLine.objects
        .filter(entry__book=book, account=account)
        .select_related("entry")
        .order_by("entry__date", "entry__entry_no", "pk")
    )

    opening = ZERO
    rows = []

    for line in lines:
        entry_date = line.entry.date

        if date_to and entry_date > date_to:
            continue

        movement = line.debit - line.credit

        if date_from and entry_date < date_from:
            opening += movement
            continue

        if date_from and not rows:
            rows.append([
                account.code,
                account.name,
                date_from,
                "",
                "",
                "Opening balance b/d",
                ZERO,
                ZERO,
                abs(opening),
                "Dr" if opening >= 0 else "Cr",
            ])

        balance = (
            opening
            if len(rows) == 1 and rows[0][5] == "Opening balance b/d"
            else ZERO
        )

        # Running balance is calculated below independently.
        rows.append([
            account.code,
            account.name,
            entry_date,
            line.entry.entry_no,
            line.entry.reference,
            line.narration or line.entry.description,
            line.debit,
            line.credit,
            ZERO,
            "",
        ])

    running = opening
    for row in rows:
        if row[5] == "Opening balance b/d":
            continue
        running += row[6] - row[7]
        row[8] = abs(running)
        row[9] = "Dr" if running >= 0 else "Cr"

    if date_from and not rows:
        rows.append([
            account.code,
            account.name,
            date_from,
            "",
            "",
            "Opening balance b/d",
            ZERO,
            ZERO,
            abs(opening),
            "Dr" if opening >= 0 else "Cr",
        ])

    return rows


def export_ledgers(
    book,
    fmt,
    account=None,
    date_from=None,
    date_to=None,
):
    accounts = (
        [account]
        if account is not None
        else list(Account.objects.order_by("code"))
    )

    sheets = []
    for current in accounts:
        rows = _ledger_rows(
            book, current, date_from, date_to
        )

        # Skip empty accounts in the all-ledgers export.
        if account is None and not rows:
            continue

        title = f"{current.pk}-{current.code}"
        

        sheets.append((
            title,
            LEDGER_HEADERS,
            rows,
            [7, 8, 9],
        ))

    if not sheets:
        sheets = [
            ("Ledgers", LEDGER_HEADERS, [], [7, 8, 9])
        ]

    filename = (
        f"book_{book.pk}_account_{account.pk}_ledger"
        if account is not None
        else f"book_{book.pk}_all_ledgers"
    )

    return _deliver(sheets, fmt, filename)
